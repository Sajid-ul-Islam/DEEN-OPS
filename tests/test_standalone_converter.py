"""
Unit tests for the standalone Product-Wise to Pathao Bulk Order Converter.
"""

import os
import tempfile
import openpyxl
import pandas as pd
import pytest

from tools.excel_pathao_converter.convert_to_pathao import (
    PATHAO_COLUMNS,
    categorize_item,
    convert_file,
    convert_orders,
    normalize_phone,
)


def test_normalize_phone():
    """Test BD phone number normalization."""
    assert normalize_phone("01711223344") == "01711223344"
    assert normalize_phone("+8801811223344") == "01811223344"
    assert normalize_phone("8801911223344") == "01911223344"
    assert normalize_phone("1711223344") == "01711223344"
    assert normalize_phone("01711-223344") == "01711223344"


def test_categorize_item():
    """Test category mapping for item description."""
    assert categorize_item("Premium Panjabi Black - L") == "Panjabi"
    assert categorize_item("Cotton Pajama White") == "Pajama"
    assert categorize_item("Drop Shoulder Oversized Tee") == "Drop Shoulder"
    assert categorize_item("Classic Polo Navy") == "Polo"
    assert categorize_item("Slim Jeans Blue") == "Jeans"
    assert categorize_item("Attar 12ml") == "Attar"


def test_convert_orders_grouping_and_columns():
    """Test converting product-wise orders into Pathao Bulk format."""
    data = [
        {
            "Order ID": "1001",
            "Customer Name": "Sajid Islam",
            "Phone": "01711223344",
            "Shipping Address": "House 1, Road 2, Banani",
            "City": "Dhaka",
            "Item Name": "Panjabi Silk",
            "Quantity": 2,
            "Total Amount": 3000,
            "Payment Method": "Cash on Delivery",
        },
        {
            "Order ID": "1001",
            "Customer Name": "Sajid Islam",
            "Phone": "01711223344",
            "Shipping Address": "House 1, Road 2, Banani",
            "City": "Dhaka",
            "Item Name": "Cotton Pajama",
            "Quantity": 1,
            "Total Amount": 3000,
            "Payment Method": "Cash on Delivery",
        },
        {
            "Order ID": "1002",
            "Customer Name": "Rahim",
            "Phone": "+8801811223344",
            "Shipping Address": "GEC Circle",
            "City": "Chattogram",
            "Item Name": "Drop Shoulder Tee",
            "Quantity": 1,
            "Total Amount": 800,
            "Payment Method": "bKash Online",
        },
    ]

    df_raw = pd.DataFrame(data)
    out_df = convert_orders(df_raw, store_name="Deen Commerce")

    assert list(out_df.columns) == PATHAO_COLUMNS
    assert len(out_df) == 2  # 3 line items grouped into 2 distinct orders

    # Order 1001 checks
    row_1001 = out_df[out_df["MerchantOrderId"] == "1001"].iloc[0]
    assert row_1001["RecipientName(*)"] == "Sajid Islam"
    assert row_1001["RecipientPhone(*)"] == "01711223344"
    assert row_1001["ItemQuantity"] == 3
    assert "2x Panjabi" in row_1001["ItemDesc"]
    assert "1x Pajama" in row_1001["ItemDesc"]
    assert row_1001["AmountToCollect(*)"] == 3000

    # Order 1002 checks (prepaid)
    row_1002 = out_df[out_df["MerchantOrderId"] == "1002"].iloc[0]
    assert row_1002["RecipientPhone(*)"] == "01811223344"
    assert row_1002["AmountToCollect(*)"] == 0
    assert row_1002["ItemQuantity"] == 1


def test_convert_file_e2e():
    """Test end-to-end file conversion with Excel input and output."""
    import gc

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir:
        input_xlsx = os.path.join(tmp_dir, "input_orders.xlsx")
        output_xlsx = os.path.join(tmp_dir, "output_pathao.xlsx")

        raw_data = [
            {
                "Order Number": "5001",
                "Full Name (Shipping)": "Test Customer",
                "Phone (Billing)": "01700112233",
                "Address 1&2 (Shipping)": "Dhanmondi 27",
                "City (Shipping)": "Dhaka",
                "Item Name": "Polo T-Shirt",
                "Quantity": 1,
                "Order Total Amount": 1200,
                "Payment Method Title": "Cash on delivery",
            }
        ]
        pd.DataFrame(raw_data).to_excel(input_xlsx, index=False)

        result_path = convert_file(input_xlsx, output_xlsx)
        assert os.path.exists(result_path)

        wb = openpyxl.load_workbook(result_path)
        try:
            assert "Pathao Bulk" in wb.sheetnames
            ws = wb["Pathao Bulk"]

            # Check headers
            headers = [cell.value for cell in ws[1]]
            assert headers == PATHAO_COLUMNS

            # Check data row
            row2 = [cell.value for cell in ws[2]]
            assert row2[0] == "Parcel"
            assert row2[1] == "Deen Commerce"
            assert row2[2] == "5001"
            assert row2[3] == "Test Customer"
            assert row2[4] == "01700112233"
            assert row2[9] == 1200
        finally:
            wb.close()
            del wb
            gc.collect()
